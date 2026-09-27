#!/usr/bin/env python3
"""Build a reviewer table from 01-TRAINING.xlsx. Does not edit the workbook or the meaning bank.

The workbook is read-only. Output is a separate review CSV plus a hash manifest.
Seed and law rows are lineage, not extra training meanings. Gold and benchmark
sections are counted and then excluded from the review queue.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

BANK = REPO / "training" / "data" / "meaning_bank.jsonl"

EXPECTED_SECTIONS = {
    "meaning_bank": 164,
    "conversation_seeds": 150,
    "law_gov": 60,
}

FIELD_MAP = {
    "english": "english",
    "devanagari": "ne_formal",
    "devanagari_informal": "ne_informal",
    "roman": "roman_formal",
    "roman_informal": "roman_informal",
}

REVIEW_COLUMNS = [
    "meaning_id",
    "workbook_row",
    "english",
    "ne_formal",
    "ne_informal",
    "roman_formal",
    "roman_informal",
    "provenance",
    "source_file",
    "source_rights",
    "register_applicability",
    "reviewer_correction_ne_formal",
    "reviewer_correction_ne_informal",
    "reviewer_correction_roman_formal",
    "reviewer_correction_roman_informal",
    "reviewer_decision",
    "reviewer",
    "reviewer_date",
    "exclusion_reason",
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def load_bank() -> dict[str, dict]:
    rows = {}
    for line in BANK.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        rows[row["meaning_id"]] = row
    return rows


def cell(row: dict, *names: str) -> str:
    for name in names:
        if name in row and row[name] is not None:
            return str(row[name]).strip()
    return ""


def read_workbook(path: Path) -> list[dict]:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("openpyxl is required to read the workbook") from exc
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    out = []
    for excel_row, values in enumerate(rows, start=2):
        record = {header[i]: values[i] if i < len(values) else None for i in range(len(header))}
        record["_workbook_row"] = excel_row
        if any(v not in (None, "") for k, v in record.items() if k != "_workbook_row"):
            out.append(record)
    book.close()
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description="Export a meaning-bank review table from the workbook")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    if not args.input.is_file():
        raise SystemExit(f"workbook not found: {args.input}")
    if not BANK.is_file():
        raise SystemExit(f"meaning bank not found: {BANK}")

    workbook_rows = read_workbook(args.input)
    bank = load_bank()
    sections = Counter(cell(row, "section", "set_name") for row in workbook_rows)
    mismatches: list[str] = []
    for name, expected in EXPECTED_SECTIONS.items():
        if sections.get(name) != expected:
            mismatches.append(f"{name}: workbook={sections.get(name, 0)} expected={expected}")

    meaning_rows = [row for row in workbook_rows if cell(row, "section", "set_name") == "meaning_bank"]
    seen_ids = []
    review = []
    for row in meaning_rows:
        meaning_id = cell(row, "id", "meaning_id")
        seen_ids.append(meaning_id)
        repo = bank.get(meaning_id)
        if repo is None:
            mismatches.append(f"workbook id missing from repo bank: {meaning_id}")
            continue
        for sheet_name, bank_name in FIELD_MAP.items():
            sheet_value = cell(row, sheet_name, bank_name)
            if sheet_value != str(repo.get(bank_name) or "").strip():
                mismatches.append(f"{meaning_id} {bank_name} differs between workbook and repo")
                break
        else:
            review.append(
                {
                    "meaning_id": meaning_id,
                    "workbook_row": row["_workbook_row"],
                    "english": repo["english"],
                    "ne_formal": repo["ne_formal"],
                    "ne_informal": repo["ne_informal"],
                    "roman_formal": repo.get("roman_formal") or "",
                    "roman_informal": repo.get("roman_informal") or "",
                    "provenance": repo.get("provenance") or "",
                    "source_file": "training/data/meaning_bank.jsonl",
                    "source_rights": "",
                    "register_applicability": "",
                    "reviewer_correction_ne_formal": "",
                    "reviewer_correction_ne_informal": "",
                    "reviewer_correction_roman_formal": "",
                    "reviewer_correction_roman_informal": "",
                    "reviewer_decision": "",
                    "reviewer": "",
                    "reviewer_date": "",
                    "exclusion_reason": "",
                }
            )

    repo_only = sorted(set(bank) - set(seen_ids))
    if repo_only:
        mismatches.append(f"repo bank ids missing from workbook: {len(repo_only)}")
    if len(set(seen_ids)) != len(seen_ids):
        mismatches.append("duplicate meaning ids in the workbook meaning_bank section")

    if mismatches:
        for item in mismatches:
            print(f"[audit] {item}", file=sys.stderr)
        raise SystemExit(f"workbook and repo are not reconciled ({len(mismatches)} checks)")

    args.out.mkdir(parents=True, exist_ok=True)
    review_path = args.out / "meaning_review.csv"
    with review_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REVIEW_COLUMNS)
        writer.writeheader()
        writer.writerows(review)

    lineage = [
        {
            "meaning_id": cell(row, "id", "meaning_id"),
            "section": cell(row, "section", "set_name"),
            "workbook_row": row["_workbook_row"],
            "role": "lineage_not_a_new_meaning",
        }
        for row in workbook_rows
        if cell(row, "section", "set_name") in {"conversation_seeds", "law_gov"}
    ]
    lineage_path = args.out / "lineage.csv"
    with lineage_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["meaning_id", "section", "workbook_row", "role"])
        writer.writeheader()
        writer.writerows(lineage)

    manifest = {
        "workbook_sha256": sha256_file(args.input),
        "meaning_bank_sha256": sha256_file(BANK),
        "review_sha256": sha256_file(review_path),
        "lineage_sha256": sha256_file(lineage_path),
        "section_counts": dict(sections),
        "review_rows": len(review),
        "note": "Reviewer decision and source rights are blank. Do not train from this file.",
    }
    manifest_path = args.out / "audit_manifest.json"
    manifest_text = json.dumps(manifest, indent=2) + "\n"
    manifest_path.write_text(manifest_text, encoding="utf-8")
    manifest["manifest_sha256"] = sha256_text(manifest_text)
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(f"[audit] wrote {review_path} rows={len(review)}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
